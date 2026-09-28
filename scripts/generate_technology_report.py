"""Build the audited technology chapter from its Markdown source.

Run with the local Python installation providing ReportLab. The project venv
provides python-docx for the editable companion and PyMuPDF for PDF inspection.
These publishing dependencies do not change the application dependency files.
"""

from __future__ import annotations

import html
import re
import sys
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, CondPageBreak, Flowable, Frame, KeepTogether, PageBreak, PageTemplate,
    Paragraph, Spacer, Table, TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SOURCE = DOCS / "Technologies_et_outils_utilises.md"
PDF = DOCS / "GreenFinance_Scorer_Technologies_et_outils_utilises.pdf"
DOCX = DOCS / "GreenFinance_Scorer_Technologies_et_outils_utilises.docx"
PREVIEW = DOCS / "technology_report_preview"

NAVY = colors.HexColor("#18354A")
GREEN = colors.HexColor("#147158")
MINT = colors.HexColor("#EDF5F1")
GREY = colors.HexColor("#596C78")
LINE = colors.HexColor("#DCE7E2")
INK = colors.HexColor("#253644")
PAGE_W, PAGE_H = A4
MARGIN = 53
WIDTH = PAGE_W - 2 * MARGIN


def register_fonts():
    root = Path("C:/Windows/Fonts")
    for name, file in [("Body", "calibri.ttf"), ("Body-Bold", "calibrib.ttf"),
                       ("Body-Italic", "calibrii.ttf"), ("Body-BoldItalic", "calibriz.ttf"),
                       ("Code", "consola.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(root / file)))
    pdfmetrics.registerFontFamily("Body", normal="Body", bold="Body-Bold",
                                  italic="Body-Italic", boldItalic="Body-BoldItalic")


def inline(text):
    text = html.escape(text)
    text = re.sub(r"`([^`]+)`", r'<font name="Code" size="8.3">\1</font>', text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", text)
    # References are unobtrusive, while keeping the repository evidence visible.
    text = re.sub(r"(\[R\d+(?:, R\d+)*\])", r'<font color="#596C78" size="9">\1</font>', text)
    return text


def plain(text):
    return text.replace("**", "").replace("`", "")


def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("#"):
            n = len(line) - len(line.lstrip("#"))
            yield ("heading", n, line[n:].strip())
            i += 1
        elif line == "<!-- architecture -->":
            yield ("diagram",)
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", x) for x in cells):
                    rows.append(cells)
                i += 1
            yield ("table", rows)
        elif line.startswith("- "):
            yield ("bullet", line[2:])
            i += 1
        else:
            chunks = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(("#", "|", "- ", "<!--")):
                chunks.append(lines[i].strip())
                i += 1
            yield ("paragraph", " ".join(chunks))


class Architecture(Flowable):
    """Vector diagram of the observed application; no external image service."""

    def __init__(self):
        super().__init__()
        self.width = WIDTH
        self.height = 242

    def draw(self):
        c = self.canv

        def box(x, y, w, h, title, desc, fill=MINT):
            c.setFillColor(fill)
            c.setStrokeColor(LINE)
            c.roundRect(x, y, w, h, 6, fill=1, stroke=1)
            c.setFillColor(NAVY)
            c.setFont("Body-Bold", 10.4)
            c.drawCentredString(x + w / 2, y + h - 18, title)
            c.setFillColor(GREY)
            c.setFont("Body", 8.6)
            for j, text in enumerate(desc):
                c.drawCentredString(x + w / 2, y + h - 32 - j * 11, text)

        def arrow(x1, y1, x2, y2):
            c.setStrokeColor(GREEN)
            c.setFillColor(GREEN)
            c.setLineWidth(1)
            c.line(x1, y1, x2, y2)
            p = c.beginPath()
            p.moveTo(x2, y2)
            if y2 < y1:
                p.lineTo(x2 - 3, y2 + 5)
                p.lineTo(x2 + 3, y2 + 5)
            else:
                p.lineTo(x2 - 5, y2 - 3)
                p.lineTo(x2 - 5, y2 + 3)
            p.close()
            c.drawPath(p, fill=1, stroke=0)

        mid = WIDTH / 2
        box(mid - 143, 190, 286, 46, "Interface React / TypeScript", ["Six espaces · navigateur · cookies de session"])
        arrow(mid, 190, mid, 170)
        c.setFont("Body", 8)
        c.setFillColor(GREY)
        c.drawString(mid + 8, 178, "HTTP · REST · JSON / fichiers")
        box(0, 100, WIDTH, 70, "API FastAPI / Python", [
            "Routes /api/v1 · permissions · services métier · scoring ESG",
            "BackgroundTasks : Docling → RapidOCR → BGE-M3 → FAISS → Gemini",
        ])
        gap = 9
        w = (WIDTH - 3 * gap) / 4
        data = [
            ("PostgreSQL", ["Données métier", "SQLModel / Psycopg"]),
            ("Redis", ["Révocation", "Tentatives de connexion"]),
            ("Stockage local", ["PDF et preuves", "Artefacts documentaires"]),
            ("API Gemini", ["Service externe", "Extraction structurée"]),
        ]
        for j, (title, desc) in enumerate(data):
            x = j * (w + gap)
            arrow(x + w / 2, 100, x + w / 2, 75)
            box(x, 9, w, 66, title, desc)


class ChapterDoc(BaseDocTemplate):
    def __init__(self, filename):
        super().__init__(str(filename), pagesize=A4, leftMargin=MARGIN,
                         rightMargin=MARGIN, topMargin=57, bottomMargin=49,
                         title="Technologies et outils utilisés — GreenFinance-Scorer",
                         author="GreenFinance-Scorer", lang="fr-FR")
        frame = Frame(MARGIN, 49, WIDTH, PAGE_H - 106, leftPadding=0,
                      rightPadding=0, topPadding=0, bottomPadding=0, id="body")
        self.addPageTemplates(PageTemplate(id="chapter", frames=frame, onPage=self.furniture))
        self.current_heading = "Technologies et outils utilisés"

    def furniture(self, c, doc):
        c.saveState()
        if doc.page == 1:
            c.setFillColor(GREEN)
            c.rect(0, PAGE_H - 14, PAGE_W, 14, stroke=0, fill=1)
        else:
            c.setFont("Body-Bold", 8)
            c.setFillColor(GREEN)
            c.drawString(MARGIN, PAGE_H - 31, "GREENFINANCE-SCORER")
            c.setFont("Body", 8)
            c.setFillColor(GREY)
            c.drawRightString(PAGE_W - MARGIN, PAGE_H - 31, "Technologies et outils utilisés")
            c.setStrokeColor(LINE)
            c.line(MARGIN, PAGE_H - 39, PAGE_W - MARGIN, PAGE_H - 39)
            c.line(MARGIN, 37, PAGE_W - MARGIN, 37)
            c.setFont("Body", 8)
            c.drawString(MARGIN, 23, "Rapport de PFE · État du projet au 12 septembre 2026")
            c.drawRightString(PAGE_W - MARGIN, 23, str(doc.page))
        c.restoreState()

    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph) and hasattr(flowable, "bookmark_key"):
            text = flowable.getPlainText()
            key = flowable.bookmark_key
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(text, key, flowable.outline_level, False)
            if flowable.outline_level == 0:
                self.notify("TOCEntry", (0, text, self.page, key))


def build_pdf(parsed):
    register_fonts()
    body = ParagraphStyle("Body", fontName="Body", fontSize=10.6, leading=15,
                          textColor=INK, alignment=TA_JUSTIFY, spaceAfter=8,
                          splitLongWords=True, allowWidows=0, allowOrphans=0)
    h1 = ParagraphStyle("Section", parent=body, fontName="Body-Bold", fontSize=21,
                        leading=25, textColor=GREEN, spaceAfter=20,
                        alignment=TA_LEFT, keepWithNext=True)
    h2 = ParagraphStyle("Subsection", parent=body, fontName="Body-Bold", fontSize=12.2,
                        leading=16, textColor=NAVY, spaceBefore=11, spaceAfter=8,
                        alignment=TA_LEFT, keepWithNext=True)
    h2_table = ParagraphStyle("TableSubsection", parent=h2, keepWithNext=False)
    small = ParagraphStyle("Small", parent=body, fontSize=9.1, leading=12.7,
                           alignment=TA_LEFT)
    cell = ParagraphStyle("Cell", parent=body, fontSize=8.8, leading=11.7,
                          spaceAfter=0, alignment=TA_LEFT)
    th = ParagraphStyle("TableHeader", parent=cell, fontName="Body-Bold", textColor=colors.white)
    ref = ParagraphStyle("Reference", parent=small, fontSize=8.7, leading=12.2, spaceAfter=8)
    cover_title = ParagraphStyle("CoverTitle", parent=h1, fontSize=36, leading=42, spaceAfter=23)
    cover_sub = ParagraphStyle("CoverSub", parent=body, fontSize=15, leading=22,
                               alignment=TA_LEFT, textColor=NAVY)
    story = [Spacer(1, 90), Paragraph("GREENFINANCE-SCORER", h2), Spacer(1, 26),
             Paragraph("Technologies<br/>et outils utilisés", cover_title),
             Paragraph("Chapitre pour le rapport<br/>de projet de fin d’études", cover_sub),
             Spacer(1, 43)]
    metadata = Table([
        [Paragraph("PÉRIMÈTRE", th), Paragraph("Code, dépendances, configurations et livrables du projet", small)],
        [Paragraph("ÉTAT EXAMINÉ", th), Paragraph("12 septembre 2026 · Version applicative déclarée 0.1.0", small)],
        [Paragraph("CONTENU", th), Paragraph("10 sections · Inventaire par couche · Versions · Références internes", small)],
    ], colWidths=[113, WIDTH - 113])
    metadata.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), GREEN), ("BACKGROUND", (1, 0), (-1, -1), MINT),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12), ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 9), ("LINEBELOW", (0, 0), (-1, -1), .5, colors.white),
    ]))
    story += [metadata, Spacer(1, 47), Paragraph(
        "Un inventaire fondé sur l’implémentation observée, avec une distinction explicite entre "
        "les fonctions disponibles, les configurations de déploiement et les éléments encore non implémentés.", small),
        PageBreak(), Paragraph("Sommaire", h1)]
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("TOC0", fontName="Body", fontSize=11.4, leading=20,
                                     textColor=NAVY, leftIndent=0, firstLineIndent=0, spaceBefore=9)]
    story += [toc, Spacer(1, 28), Paragraph(
        "Les versions verrouillées sont distinguées des étiquettes d’images et des versions observées "
        "sur le poste. Les références [R01] à [R24] renvoient aux fichiers du projet.", small)]
    started = False
    n = 0
    for idx, block in enumerate(parsed):
        kind = block[0]
        if kind == "heading" and block[1] == 2:
            started = True
            story.append(PageBreak())
            p = Paragraph(inline(block[2]), h1)
            p.bookmark_key = f"section-{n}"
            p.outline_level = 0
            n += 1
            story.append(p)
        elif not started:
            continue
        elif kind == "heading":
            if block[2].startswith("Références internes"):
                story.append(PageBreak())
                style = h1
                level = 0
            else:
                style = h2
                level = 1
                if idx + 1 < len(parsed) and parsed[idx + 1][0] == "table":
                    style = h2_table
                    story.append(CondPageBreak(150))
            p = Paragraph(inline(block[2]), style)
            p.bookmark_key = f"section-{n}"
            p.outline_level = level
            n += 1
            story.append(p)
        elif kind == "paragraph":
            story.append(Paragraph(inline(block[1]), body))
        elif kind == "bullet":
            if block[1].startswith("**[R13]"):
                story.append(PageBreak())
                story.append(Paragraph("Références internes — suite", h2))
            story.append(Paragraph(inline(block[1]), ref))
        elif kind == "diagram":
            story.append(KeepTogether([Architecture(), Spacer(1, 5), Paragraph(
                "Figure 1 — Architecture logique constatée dans le code du projet.", small)]))
        elif kind == "table":
            rows = block[1]
            ratios = {2: [.24, .76], 3: [.19, .37, .44], 4: [.235, .19, .235, .34]}[len(rows[0])]
            data = [[Paragraph(inline(t), th if i == 0 else cell) for t in row] for i, row in enumerate(rows)]
            table = Table(data, colWidths=[WIDTH * r for r in ratios], repeatRows=1,
                          hAlign="LEFT", spaceBefore=6, spaceAfter=13)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), GREEN),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, MINT]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LINEBELOW", (0, 0), (-1, 0), 1, GREEN),
                ("LINEBELOW", (0, 1), (-1, -1), .35, LINE),
            ]))
            story.append(table)
    doc = ChapterDoc(PDF)
    doc.multiBuild(story)
    print(f"PDF: {PDF}")


def add_word_inline(p, text):
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not part:
            continue
        run = p.add_run(plain(part))
        if part.startswith("**"):
            run.bold = True
        elif part.startswith("`"):
            from docx.shared import Pt
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)


def build_docx(parsed):
    # Read the existing local environment; do not install or alter project dependencies.
    sys.path.append(str(ROOT / ".venv" / "Lib" / "site-packages"))
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.9)
    sec.left_margin = sec.right_margin = Cm(1.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(7)
    for level, size in [(1, 23), (2, 14)]:
        s = doc.styles[f"Heading {level}"]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string("147158" if level == 1 else "18354A")
        s.paragraph_format.keep_with_next = True
    sec.header.paragraphs[0].text = "GREENFINANCE-SCORER  |  Technologies et outils utilisés"
    sec.header.paragraphs[0].style = "Caption"
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run("Rapport de PFE  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    foot._p.append(field)
    started = False
    for block in parsed:
        kind = block[0]
        if kind == "heading":
            lev, text = block[1:]
            if lev == 1:
                doc.add_heading(text, 0)
            else:
                if lev == 2:
                    if started:
                        doc.add_page_break()
                    started = True
                doc.add_heading(text, 1 if lev == 2 else 2)
        elif kind in ("paragraph", "bullet"):
            p = doc.add_paragraph()
            if kind == "paragraph" and started:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_word_inline(p, block[1])
        elif kind == "table":
            rows = block[1]
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for i, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    add_word_inline(cell.paragraphs[0], value)
                    for r in cell.paragraphs[0].runs:
                        r.font.size = Pt(9)
                        if i == 0:
                            r.bold = True
                if i == 0:
                    repeat = OxmlElement("w:tblHeader")
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
        elif kind == "diagram":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run("React / TypeScript\n↓ HTTP · REST · JSON / fichiers\n"
                      "FastAPI / Python · services métier · pipeline documentaire\n"
                      "↓\nPostgreSQL   |   Redis   |   Stockage local   |   API Gemini").bold = True
    doc.core_properties.title = "Technologies et outils utilisés — GreenFinance-Scorer"
    doc.core_properties.subject = "Chapitre PFE fondé sur l’audit du code au 12 septembre 2026"
    doc.core_properties.author = "GreenFinance-Scorer"
    doc.save(DOCX)
    print(f"DOCX: {DOCX}")


def inspect_pdf():
    sys.path.append(str(ROOT / ".venv" / "Lib" / "site-packages"))
    import pymupdf
    import json
    document = pymupdf.open(PDF)
    PREVIEW.mkdir(exist_ok=True)
    texts = [p.get_text() for p in document]
    faults = []
    for i, page in enumerate(document):
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    x0, y0, x1, y1 = span["bbox"]
                    if x0 < 15 or y0 < 0 or x1 > PAGE_W - 15 or y1 > PAGE_H:
                        faults.append({"page": i + 1, "text": span["text"], "bbox": span["bbox"]})
    selected = sorted(set([0, 1, 2, 3, len(document) - 1] +
                          [i for i, t in enumerate(texts) if any(s in t for s in
                           ["IndexFlatIP", "6.1. Organisation", "10.1. Frontend", "10.3. IA", "10.4. Déploiement"])]))
    thumbs = []
    for i in selected:
        pix = document[i].get_pixmap(matrix=pymupdf.Matrix(1.15, 1.15), alpha=False)
        out = PREVIEW / f"page_{i+1:02d}.png"
        pix.save(out)
        thumbs.append((i + 1, out))
    from PIL import Image, ImageDraw, ImageFont
    thumb_w, thumb_h = 298, 421
    cols = 3
    rows = (len(thumbs) + cols - 1) // cols
    contact = Image.new("RGB", (cols * (thumb_w + 18) + 18, rows * (thumb_h + 38) + 18), "#E5ECE8")
    draw = ImageDraw.Draw(contact)
    font = ImageFont.truetype("C:/Windows/Fonts/calibri.ttf", 15)
    for j, (page_no, path) in enumerate(thumbs):
        im = Image.open(path).convert("RGB")
        im.thumbnail((thumb_w, thumb_h))
        x, y = 18 + (j % cols) * (thumb_w + 18), 18 + (j // cols) * (thumb_h + 38)
        contact.paste(im, (x, y))
        draw.text((x, y + thumb_h + 5), f"Page {page_no}", fill="#18354A", font=font)
    contact.save(PREVIEW / "contact_sheet.jpg", quality=90)
    report = {
        "pages": len(document), "pdf_bytes": PDF.stat().st_size,
        "docx_bytes": DOCX.stat().st_size,
        "words_markdown": len(SOURCE.read_text(encoding="utf-8").split()),
        "outline_entries": len(document.get_toc()), "page_text_lengths": [len(t) for t in texts],
        "out_of_page_text": faults, "preview_pages": [x[0] for x in thumbs],
        "missing_major_sections": [f"{i}." for i in range(1, 11)
                                   if not any(re.search(rf"\b{i}\.\s", t) for t in texts)],
        "replacement_characters": sum(t.count(chr(65533)) for t in texts),
    }
    (PREVIEW / "validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    content = list(blocks(SOURCE.read_text(encoding="utf-8")))
    build_pdf(content)
    build_docx(content)
    inspect_pdf()
